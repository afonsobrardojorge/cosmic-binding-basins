from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT = Path(__file__).resolve().parents[1]
DOCX_OUT = OUT / "paper" / "Cosmic_Binding_Basins_preprint_v0.2.docx"
CSV = OUT / "data" / "tec_literature_validation.csv"
SUMMARY_JSON = OUT / "data" / "tec_literature_validation_summary.json"
FIG_PHASE = OUT / "figures" / "tec_diagrama_fase_Tlambda.png"
FIG_SUMMARY = OUT / "figures" / "tec_resumo_por_classe.png"
FIG_VALIDATION = OUT / "figures" / "tec_literature_validation.png"
FIG_TURNAROUND = OUT / "figures" / "tec_turnaround_fit.png"
FIG_DYNAMIC_DE = OUT / "figures" / "tec_energia_escura_dinamica.png"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D5DBE3", size="8") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=85, start=95, bottom=85, end=95) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def mark_header_row(table) -> None:
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)


def style_table(table, header_fill="203B56") -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    mark_header_row(table)
    for row_i, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_border(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8)
            if row_i == 0:
                set_cell_shading(cell, header_fill)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            elif row_i % 2 == 0:
                set_cell_shading(cell, "F5F7FA")


def add_equation(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.name = "Cambria Math"
    r.font.size = Pt(11)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(85, 85, 85)


def setup_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.85)
    section.right_margin = Cm(1.85)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    styles["Normal"].paragraph_format.line_spacing = 1.04
    for name, size, color in [
        ("Heading 1", 13, RGBColor(31, 62, 88)),
        ("Heading 2", 10.5, RGBColor(55, 85, 110)),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.bold = True
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(4)


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Cosmic Binding Basins")
    r.font.name = "Arial"
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = RGBColor(31, 62, 88)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("A Dimensionless Boundary Index for Gravitational Membership in ΛCDM")
    r2.font.name = "Arial"
    r2.font.size = Pt(12)
    r2.italic = True
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run("Short exploratory preprint draft v0.2 | 20 May 2026 | Afonso Brardo Buinheira Sal Jorge")
    r3.font.size = Pt(8.5)
    r3.font.color.rgb = RGBColor(90, 90, 90)


def add_object_table(doc: Document, df: pd.DataFrame) -> None:
    radius_type = {
        "satélite ligado, não turnaround": "bound satellite (not turnaround)",
        "zero-velocity / turnaround local": "zero-velocity / local turnaround",
        "zero-velocity / turnaround aproximado": "zero-velocity / approximate turnaround",
        "raio interno, não turnaround": "internal core radius (not turnaround)",
        "turnaround/infall bound; ambiente filamentar": "turnaround/infall bound; filamentary environment",
        "turnaround/collapsing core": "turnaround/collapsing core",
        "turnaround/candidate bound supercluster": "turnaround/candidate bound supercluster",
    }
    table = doc.add_table(rows=1, cols=7)
    hdr = table.rows[0].cells
    headers = ["Object", "Radius type", "M [M_sun]", "R [Mpc]", "TΛ", "Uncertainty flag", "Primary source"]
    for i, h in enumerate(headers):
        hdr[i].text = h
    for _, row in df.iterrows():
        cells = table.add_row().cells
        cells[0].text = row["estrutura"]
        cells[1].text = radius_type.get(row["tipo_medida"], row["tipo_medida"])
        cells[2].text = f"{row['massa_msun']:.2e}"
        cells[3].text = f"{row['raio_mpc']:.2f}"
        cells[4].text = f"{row['T_lambda']:.3f}"
        cells[5].text = "possible" if row["pode_violar_com_incerteza"] else "no"
        cells[6].text = str(row["fonte"]).split(";")[0]
    style_table(table, header_fill="1F4E79")


def write_doc() -> None:
    df = pd.read_csv(CSV)
    summary = json.loads(SUMMARY_JSON.read_text(encoding="utf-8"))
    doc = Document()
    setup_doc(doc)
    doc.core_properties.title = "Cosmic Binding Basins short preprint"
    doc.core_properties.author = "Afonso Brardo Buinheira Sal Jorge"

    add_title(doc)

    doc.add_heading("Abstract", level=1)
    doc.add_paragraph(
        "We propose Cosmic Binding Basins (TEC-2), a phenomenological ΛCDM-compatible framework in which dark energy is treated as a large-scale boundary-setting term for gravitational membership, rather than as an internal dynamical disruption mechanism within already bound systems. The central diagnostic is a dimensionless local-global index comparing the effective acceleration associated with the cosmological constant to the mean gravitational acceleration of an enclosed mass."
    )
    add_equation(doc, "TΛ(R,M) = ΩΛ H0² R³ / (GM)")
    doc.add_paragraph(
        "This work is not claimed as a new physical law, but as a compact diagnostic derived from the ΛCDM turnaround scale."
    )
    doc.add_paragraph(
        f"We test the index in a pilot literature sample of {summary['n_total']} published structures, including Local Group, M81/M82, Fornax-Eridanus, Virgo, Coma, and several collapsing supercluster cores. Among {summary['n_informative_turnaround_like']} turnaround-like or zero-velocity entries, all central estimates satisfy TΛ ≤ 1, with median TΛ = {summary['median_T_informative']:.3f}. This does not prove universality, but provides preliminary independent consistency with TEC-2 as an exterior-boundary criterion."
    )

    doc.add_heading("1. Physical Idea", level=1)
    doc.add_paragraph(
        "The guiding claim is deliberately limited: TEC-2 does not introduce a new force. It reorganizes a known ΛCDM boundary scale into an operational index. The phrase 'structural tension' refers to the balance between local gravitational aggregation and global accelerated expansion. In this view, galaxies, halos, clusters and superclusters are nested gravitational basins; black holes are extreme inner basins, not exceptions to the framework."
    )
    doc.add_paragraph(
        "The proposed publication-level statement is: dark energy is treated here as a large-scale boundary-setting term for gravitational membership, rather than as an internal dynamical disruption mechanism within already bound systems. The testable question is whether observed boundaries such as zero-velocity or turnaround surfaces lie at or below TΛ ≈ 1."
    )

    doc.add_heading("2. Formal Derivation Of The Base Index", level=1)
    doc.add_paragraph(
        "The base index follows from comparing two accelerations. For a mass M enclosed within physical radius R, the characteristic inward gravitational acceleration is"
    )
    add_equation(doc, "ag(R,M) = GM/R²")
    doc.add_paragraph(
        "For a cosmological constant in the low-redshift ΛCDM limit, the outward acceleration scale associated with dark energy is"
    )
    add_equation(doc, "aΛ(R) ≈ ΩΛ H0² R")
    doc.add_paragraph(
        "The boundary regime is obtained when aΛ and ag are comparable. Their ratio defines a dimensionless continuous diagnostic:"
    )
    add_equation(doc, "TΛ(R,M) = aΛ/ag = ΩΛ H0² R³/(GM)")
    doc.add_paragraph(
        "Thus TΛ < 1 indicates that local gravity dominates the analysed scale; TΛ ≈ 1 marks the exterior boundary regime; TΛ > 1 indicates that the analysed scale lies beyond the gravitational basin in the ideal isolated limit."
    )

    doc.add_heading("3. Relation To Existing ΛCDM Boundaries", level=1)
    doc.add_paragraph(
        "For a cosmological constant, TΛ = 1 is equivalent to the maximum turnaround-radius scale in Schwarzschild-de Sitter/Kottler and spherical-collapse arguments:"
    )
    add_equation(doc, "Rta,max = (3GM/Λc²)^(1/3) = [GM/(ΩΛH0²)]^(1/3)")
    doc.add_paragraph(
        "Thus TEC-2 should not be presented as a replacement for turnaround theory, virial radii, spherical collapse or ΛCDM. Its proposed contribution is classification: a continuous dimensionless index applicable across inner cores, virialized regions, turnaround-like surfaces, and large-scale structures."
    )
    table = doc.add_table(rows=1, cols=3)
    hdr = table.rows[0].cells
    hdr[0].text = "Known framework"
    hdr[1].text = "What is already known"
    hdr[2].text = "TEC-2 contribution"
    rows = [
        ("ΛCDM turnaround", "There is a maximum mass-dependent exterior scale for bound structures.", "Recasts the scale as a continuous membership index TΛ."),
        ("Spherical collapse", "Ideal overdensities pass through expansion, turnaround, collapse and virialization.", "Uses the boundary idea as a classifier across object classes."),
        ("Virial radius", "Describes an internal approximately relaxed halo region.", "Separates internal relaxation from exterior gravitational-basin membership."),
        ("Cosmic-web classifiers", "Classify knots, filaments, sheets and voids by density/tidal geometry.", "Adds an acceleration-ratio criterion tied to dark-energy dominance."),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    style_table(table, header_fill="385723")

    doc.add_heading("4. Extensions: Dynamic Dark Energy, Environment And Geometry", level=1)
    doc.add_paragraph(
        "The natural extension is not to invent a new force, but to generalize the acceleration ratio. For a dark-energy equation of state w(z), the repulsive contribution in the acceleration equation is proportional to -[1 + 3w(z)]/2. A dynamical dark-energy version of the index is therefore"
    )
    add_equation(doc, "TDE(R,M,z) = {-[1 + 3w(z)]/2} ΩDE(z) H(z)² R³/(GM)")
    doc.add_paragraph(
        "For w = -1 and z = 0, this reduces to TΛ. This version asks a sharper question: if dark energy evolves with time, do the boundaries of gravitational basins evolve measurably with redshift?"
    )
    if FIG_DYNAMIC_DE.exists():
        fig = doc.add_picture(str(FIG_DYNAMIC_DE), width=Inches(6.4))
        fig._inline.docPr.set("descr", "Dynamical dark energy figure showing how w(z) can shift the TEC boundary term with redshift.")
        fig._inline.docPr.set("title", "Dynamical dark energy extension")
        add_caption(doc, "Figure 1. TEC extension for dynamical dark energy. Curves illustrate how different w(z) histories could shift the boundary-setting term with redshift.")
    doc.add_paragraph(
        "A later, more realistic version should include environmental corrections, because real structures are not isolated spheres. We propose the following future extension, not yet validated in this preprint:"
    )
    add_equation(doc, "Tbasin(R,M,z,η) = TDE(R,Meff,z) · Fenv(η)")
    doc.add_paragraph(
        "Here Meff(R) is the lensing or dynamical mass enclosed by the basin, and Fenv(η) is an environmental factor describing isolation, filamentarity, tidal field or proximity to neighbouring masses. This is the most original future direction, but it must be fitted from simulations or survey data before being treated as physical."
    )
    doc.add_heading("4.1 Geometric Basin Dominance (TEC-3 Proposal)", level=2)
    doc.add_paragraph(
        "A stronger conceptual extension is to express basin membership geometrically. In general relativity, local masses generate Weyl curvature, while the cosmological constant sets an approximately uniform background curvature scale. TEC-3 is therefore proposed as a geometric reformulation: a region belongs to a gravitational basin while local Weyl or tidal curvature dominates over the dark-energy curvature scale."
    )
    add_equation(doc, "BΛ(x) = sqrt(Cαβγδ C^αβγδ / 48) / (Λ/3)")
    doc.add_paragraph(
        "In the Schwarzschild-de Sitter limit, sqrt(Cαβγδ C^αβγδ/48) = GM/(c²R³) and Λ/3 = ΩΛH0²/c², so"
    )
    add_equation(doc, "BΛ(R) = GM/(ΩΛ H0² R³) = 1/TΛ(R,M)")
    doc.add_paragraph(
        "Thus the original acceleration index is recovered as the inverse of the geometric dominance index in the ideal spherical case. The interpretation is BΛ > 1 for local-curvature dominance, BΛ ≈ 1 at the geometric basin boundary, and BΛ < 1 when the cosmological background curvature dominates."
    )
    doc.add_paragraph(
        "For simulations, a more practical proxy is the traceless tidal tensor of the gravitational potential, Eij = ∂i∂jΦ - δij∇²Φ/3. A testable tidal form is"
    )
    add_equation(doc, "BE(x,z) = ||Eij(x,z)|| / ADE(z)")
    add_equation(doc, "ADE(z) = -[1 + 3w(z)] ΩDE(z) H(z)² / 2")
    doc.add_paragraph(
        "This TEC-3 section is a proposed research direction, not a validated result of the present paper. Its value is that it connects the mass-radius diagnostic to relativistic curvature and to quantities that can be measured in N-body simulations."
    )

    doc.add_heading("5. Pilot Data And Radius Definitions", level=1)
    doc.add_paragraph(
        "The pilot sample is intentionally small and heterogeneous. Each entry is labelled by radius type, because mixing virial radius, zero-velocity radius, turnaround radius and dense collapsing core radius would otherwise overstate the evidence. Core or virial points are controls; turnaround-like and zero-velocity points are the informative boundary tests."
    )
    if FIG_PHASE.exists():
        fig = doc.add_picture(str(FIG_PHASE), width=Inches(6.4))
        fig._inline.docPr.set("descr", "TEC phase diagram showing synthetic object classes relative to the T_lambda equals one boundary.")
        fig._inline.docPr.set("title", "TEC phase diagram")
        add_caption(doc, "Figure 2. TEC phase diagram: mass, scale and TΛ regime in a synthetic proof-of-concept catalogue.")
    if FIG_SUMMARY.exists():
        fig = doc.add_picture(str(FIG_SUMMARY), width=Inches(6.4))
        fig._inline.docPr.set("descr", "Median log10 T_lambda by synthetic class.")
        fig._inline.docPr.set("title", "Synthetic class summary")
        add_caption(doc, "Figure 3. Median log10(TΛ) by synthetic class. Negative values indicate bound-regime scales; positive values indicate expansion-dominated scales.")
    add_object_table(doc, df)

    doc.add_heading("6. Results", level=1)
    doc.add_paragraph(
        f"All {summary['n_informative_turnaround_like']} central turnaround-like estimates satisfy TΛ ≤ 1. The largest central value is TΛ = {summary['max_T_central_informative']:.3f}. Coma is the only object whose uncertainty envelope can approach or exceed the boundary, consistent with its known filamentary, non-isolated environment and uncertainty in the outer infall/turnaround scale."
    )
    if FIG_VALIDATION.exists():
        fig = doc.add_picture(str(FIG_VALIDATION), width=Inches(6.4))
        fig._inline.docPr.set("descr", "Pilot literature validation plot with mass-radius error bars and T_lambda equals one boundary.")
        fig._inline.docPr.set("title", "Pilot literature validation")
        add_caption(doc, "Figure 4. Published structures in the mass-radius plane. Error bars show reported or adopted mass/radius uncertainty ranges. The solid black line is the TΛ = 1 boundary, i.e. R = RΛ(M).")

    if FIG_TURNAROUND.exists():
        fig = doc.add_picture(str(FIG_TURNAROUND), width=Inches(6.4))
        fig._inline.docPr.set("descr", "Synthetic regression showing R proportional to M to one third for turnaround-like boundaries.")
        fig._inline.docPr.set("title", "Synthetic turnaround scaling")
        add_caption(doc, "Figure 5. Synthetic sanity check recovering the expected R ∝ M^(1/3) scaling. This is a methodological check, not observational proof.")

    doc.add_heading("7. Discussion", level=1)
    doc.add_paragraph(
        "The pilot result supports a cautious claim: structures already interpreted in the literature as bound, in turnaround, or in collapse are consistent with the TEC-2 boundary criterion. The result is not yet a discovery of new physics. It is a compact re-expression of a known ΛCDM scale that appears useful as a gravitational-basin classifier. The TEC-3 geometric extension gives the framework a deeper relativistic form, but remains a proposal until tested with simulations."
    )
    doc.add_paragraph(
        "The next decisive test is comparative. TΛ must be evaluated against existing metrics: mass alone, mean density, virial radius, turnaround radius, phase-space caustics, and cosmic-web classifiers. A publishable follow-up should test whether TΛ improves classification or reduces scatter in boundary prediction using a larger homogeneous simulation or survey sample."
    )

    doc.add_heading("8. Limitations", level=1)
    doc.add_paragraph(
        "The present sample is small, heterogeneous, and partly based on different definitions of boundary radius. Some masses and radii are model dependent, and supercluster cores are not spherical isolated systems. The current analysis therefore provides preliminary independent consistency, not universal proof."
    )

    doc.add_heading("9. Conclusion", level=1)
    doc.add_paragraph(
        "The validation pilot shows that published structures interpreted as bound, in turnaround, or collapsing remain below TΛ = 1 in their central estimates, consistent with TEC-2 as an outer-boundary criterion for gravitational membership. This result does not demonstrate universality, but it gives the hypothesis a defensible, testable preprint-level form."
    )

    doc.add_heading("References", level=1)
    references = [
        "Aghanim et al. (Planck Collaboration), 2020, A&A, 641, A6.",
        "Pavlidou & Tomaras, 2014, JCAP, 2014(09), 020, arXiv:1310.1920.",
        "Karachentsev & Kashibadze, 2006, Astrophysics, 49, 3.",
        "Nasonova, de Freitas Pacheco & Karachentsev, 2011, A&A, 532, A104.",
        "Benisty, Wagner, Haridasu & Salucci, 2025, arXiv:2504.04135.",
        "Einasto et al., 2021, A&A, 649, A51, arXiv:2103.02326.",
        "Reisenegger et al., 2000, AJ, 120, 523.",
        "Kormendy & Ho, 2013, ARA&A, 51, 511.",
        "Binney & Tremaine, 2008, Galactic Dynamics, 2nd ed.",
    ]
    for ref in references:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.45)
        p.paragraph_format.first_line_indent = Cm(-0.45)
        p.add_run(ref)

    doc.save(DOCX_OUT)
    print(DOCX_OUT)


if __name__ == "__main__":
    write_doc()
